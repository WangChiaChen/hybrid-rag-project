"""把截圖插進「比賽簡報_升級版.pptx」的截圖佔位——因為 Office 未啟動不能在 PowerPoint 裡編輯，
改用程式插，繞過編輯鎖。

## 用法

1. 把截圖存成 PNG／JPG，放進 簡報與文件/screenshots/，檔名照下面 SLOTS 的 key：
     s16_dashboard.png   分析儀表板實機截圖
     s17_compare.png     跨機構比較實機截圖
     s18_qa.png          問答與報告實機截圖（例：國泰各子公司稅後淨利）
     s07_graph.png       EAP『探索關連』知識圖譜截圖
     s11_settings.png    EAP 後台 Robot Settings／Context 截圖
     s15a_eap.png        EAP 平台問答截圖（國泰世華 132 億、附出處）
     s15b_crosscheck.png EAP 答＋本地紅框標出換算差
   （有幾張放幾張，沒有的自動跳過，之後補了再跑一次即可。）

2. 執行：
     venv/Scripts/python.exe scripts/insert_ppt_screenshots.py

每張圖會等比例縮放、置中貼進對應投影片的圖框，並清掉該處的「（貼截圖…）」提示字。
原檔就地更新（已 git 追蹤，隨時可 git checkout 還原）。
"""
import os
import sys

sys.stdout.reconfigure(encoding="utf-8")

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DECK = os.path.join(BASE, "簡報與文件", "比賽簡報_升級版.pptx")
SHOTS = os.path.join(BASE, "簡報與文件", "screenshots")

# key（檔名去副檔名） -> (投影片序號1起, 左, 上, 寬, 高 吋, 要清掉提示字的 shape 索引清單)
SLOTS = {
    "s07_graph":       (7,  0.7, 2.2, 7.4, 3.9, [5]),
    "s11_settings":    (11, 0.7, 2.2, 12.0, 3.0, []),
    "s15a_eap":        (15, 0.7, 2.6, 6.0, 3.5, [6]),
    "s15b_crosscheck": (15, 7.0, 2.6, 5.6, 3.5, [9]),
    "s16_dashboard":   (16, 6.5, 2.2, 6.1, 4.2, [9]),
    "s17_compare":     (17, 6.5, 2.2, 6.1, 4.2, [9]),
    "s18_qa":          (18, 6.5, 2.2, 6.1, 4.2, [9]),
}
EXTS = (".png", ".jpg", ".jpeg")


def _find(key):
    for ext in EXTS:
        p = os.path.join(SHOTS, key + ext)
        if os.path.exists(p):
            return p
    return None


def _clear_hint(shape):
    """清掉「（貼截圖…）」那行提示字，保留其他文字（如區塊標題）。"""
    if not shape.has_text_frame:
        return
    for para in shape.text_frame.paragraphs:
        joined = "".join(r.text for r in para.runs)
        if "貼" in joined and ("截圖" in joined or "Demo" in joined):
            for r in para.runs:
                r.text = ""


def main():
    os.makedirs(SHOTS, exist_ok=True)
    if not os.path.exists(DECK):
        print("找不到簡報：", DECK)
        return
    prs = Presentation(DECK)
    done, missing = [], []
    for key, (idx, left, top, w, h, hint_shapes) in SLOTS.items():
        img = _find(key)
        if not img:
            missing.append(key)
            continue
        slide = prs.slides[idx - 1]
        # 等比例縮放、置中貼進圖框
        iw, ih = Image.open(img).size
        scale = min(w / (iw / 96.0), h / (ih / 96.0))   # 以 96dpi 推吋，取能塞進框的比例
        pw, ph = (iw / 96.0) * scale, (ih / 96.0) * scale
        px, py = left + (w - pw) / 2, top + (h - ph) / 2
        slide.shapes.add_picture(img, Inches(px), Inches(py), Inches(pw), Inches(ph))
        for si in hint_shapes:
            if si < len(slide.shapes):
                _clear_hint(slide.shapes[si])
        done.append((key, idx))

    if done:
        prs.save(DECK)
    print(f"已插入 {len(done)} 張截圖：")
    for key, idx in done:
        print(f"  ✓ 投影片{idx}  ← {key}")
    if missing:
        print(f"\n還缺 {len(missing)} 張（放進 {SHOTS} 再跑一次）：")
        for key in missing:
            print(f"  - {key}（第 {SLOTS[key][0]} 頁）")
    if not done:
        print("（screenshots 資料夾裡還沒有可用的圖，什麼都沒改）")


if __name__ == "__main__":
    main()
